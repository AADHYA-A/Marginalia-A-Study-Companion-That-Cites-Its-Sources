"""Marginalia backend: ingest lecture/textbook/slides, cited Q&A, adaptive quizzes."""

import os
import io
import re
import json
import random
import time
import base64
import tempfile
from pathlib import Path

import numpy as np
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.cluster import KMeans
from sklearn.metrics.pairwise import cosine_similarity

ROOT = Path(__file__).parent
# Vercel's filesystem is read-only except for /tmp.
DATA = Path("/tmp/marginalia") if os.getenv("VERCEL") else ROOT / "data"
try:
    DATA.mkdir(parents=True, exist_ok=True)
except OSError:
    DATA = Path("/tmp")
STORE = DATA / "store.json"
MODEL = os.getenv("CLAUDE_MODEL", "claude-sonnet-5-5")

client = None
if os.getenv("ANTHROPIC_API_KEY"):
    import anthropic
    client = anthropic.Anthropic()

S = {"chunks": [], "mastery": {}, "history": []}
try:
    if STORE.exists():
        S = json.loads(STORE.read_text(encoding="utf-8"))
except Exception:
    pass

vec = None
mat = None
app = FastAPI(title="Marginalia")

@app.get("/")
def home():
    index_file = ROOT / "static" / "index.html"
    if index_file.exists():
        return FileResponse(index_file)
    return {
        "message": "Marginalia API is running",
        "docs": "/docs",
        "state": "/api/state"
    }

def save():
    try:
        STORE.write_text(json.dumps(S), encoding="utf-8")
    except OSError:
        # Continue in memory if the filesystem is not writable.
        pass


def fmt(seconds):
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def pub(chunk):
    return {key: chunk[key] for key in ("id", "type", "ref", "source", "text", "concept") if key in chunk}


def llm(prompt, system="", max_tokens=700):
    if client is None:
        raise RuntimeError("ANTHROPIC_API_KEY is not configured")
    response = client.messages.create(
        model=MODEL,
        max_tokens=max_tokens,
        system=system,
        messages=[{"role": "user", "content": prompt}],
    )
    return response.content[0].text


# ---------- indexing ----------
def reindex():
    global vec, mat
    chunks = S["chunks"]
    if not chunks:
        vec = mat = None
        S["mastery"] = {}
        return

    vec = TfidfVectorizer(stop_words="english", sublinear_tf=True, ngram_range=(1, 2))
    mat = vec.fit_transform([chunk["text"] for chunk in chunks])
    k = max(1, min(8, len(chunks) // 3))
    labels = KMeans(n_clusters=k, n_init=5, random_state=0).fit_predict(mat) if k > 1 else [0] * len(chunks)
    terms = np.array(vec.get_feature_names_out())
    names = {}
    for group in set(labels):
        indices = [i for i, label in enumerate(labels) if label == group]
        top = np.asarray(mat[indices].mean(axis=0)).ravel().argsort()[::-1][:2]
        names[group] = " / ".join(terms[top])
    for chunk, label in zip(chunks, labels):
        chunk["concept"] = names[label]
    concepts = set(names.values())
    S["mastery"] = {concept: S["mastery"].get(concept, 0.3) for concept in concepts}


def retrieve(query, k=4):
    if vec is None or mat is None:
        return []
    scores = cosine_similarity(vec.transform([query]), mat).ravel()
    return [S["chunks"][i] for i in scores.argsort()[::-1][:k] if scores[i] > 0.02]


# ---------- parsers ----------
def chunk_text(text, size=650):
    sentences = re.split(r"(?<=[.!?])\s+", re.sub(r"\s+", " ", text).strip())
    output, buffer = [], ""
    for sentence in sentences:
        if buffer and len(buffer) + len(sentence) > size:
            output.append(buffer)
            buffer = ""
        buffer = (buffer + " " + sentence).strip()
    if buffer:
        output.append(buffer)
    return output


def parse_pdf(data):
    from pypdf import PdfReader
    pages = PdfReader(io.BytesIO(data)).pages
    return [(f"p.{i + 1}", text) for i, page in enumerate(pages) for text in chunk_text(page.extract_text() or "")], "Textbook"


def caption(blob, media_type):
    """Describe a slide image with Claude Vision when an API key is configured."""
    if client is None:
        return ""
    try:
        response = client.messages.create(
            model=MODEL,
            max_tokens=200,
            messages=[{"role": "user", "content": [
                {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": base64.b64encode(blob).decode()}},
                {"type": "text", "text": "Describe this lecture slide figure in 1-2 factual sentences, including labels and numbers."},
            ]}],
        )
        return response.content[0].text
    except Exception:
        return ""


def parse_pptx(data):
    from pptx import Presentation
    output = []
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if client and shape.shape_type == 13:
                description = caption(shape.image.blob, shape.image.content_type)
                if description:
                    texts.append("[Figure] " + description)
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
        combined = " ".join(text for text in texts if text.strip())
        if combined.strip():
            output.append((f"slide {i}", combined))
    return output, "Slide"


# Supports SRT/VTT timestamps such as 00:01:02,000 --> 00:01:05,000.
TS = re.compile(r"(?:(\d+):)?(\d{1,2}):(\d{2})[,.]\d+\s*-->")


def group_segments(segments, window=45):
    if not segments:
        return []
    output, start, buffer = [], segments[0][0], ""
    for seconds, line in segments:
        if buffer and seconds - start >= window:
            output.append((fmt(start), buffer))
            start, buffer = seconds, ""
        buffer = (buffer + " " + line).strip()
    if buffer:
        output.append((fmt(start), buffer))
    return output


def parse_lecture(text):
    segments, current = [], None
    for line in text.splitlines():
        match = TS.search(line)
        if match:
            current = int(match.group(1) or 0) * 3600 + int(match.group(2)) * 60 + int(match.group(3))
            continue
        line = line.strip()
        if line and current is not None and not line.isdigit() and line != "WEBVTT":
            segments.append((current, line))
    if segments:
        return group_segments(segments), "Lecture"
    return [(f"para {i + 1}", text) for i, text in enumerate(chunk_text(text))], "Notes"


def parse_media(data, name):
    try:
        from faster_whisper import WhisperModel
    except ImportError:
        raise HTTPException(400, "Audio/video needs faster-whisper and ffmpeg installed, or upload an .srt/.vtt transcript.")
    with tempfile.NamedTemporaryFile(suffix=Path(name).suffix) as temp_file:
        temp_file.write(data)
        temp_file.flush()
        segments, _ = WhisperModel("base", compute_type="int8").transcribe(temp_file.name)
        segments = [(int(segment.start), segment.text.strip()) for segment in segments if segment.text.strip()]
    if not segments:
        raise HTTPException(400, "No speech found in that file.")
    return group_segments(segments), "Lecture"


MEDIA = {".mp3", ".wav", ".m4a", ".mp4", ".mov", ".webm", ".mkv"}


def add(items, item_type, source):
    next_id = max([chunk["id"] for chunk in S["chunks"]], default=0) + 1
    for reference, text in items:
        S["chunks"].append({"id": next_id, "type": item_type, "ref": reference, "source": source, "text": text})
        next_id += 1


# ---------- API ----------
@app.get("/api/state")
def state():
    mastery = S["mastery"]
    weakest = sorted(mastery, key=mastery.get)[:3]
    plan = [
        {"concept": concept, "mastery": mastery[concept],
         "sources": [pub(chunk) for chunk in S["chunks"] if chunk.get("concept") == concept][:2]}
        for concept in weakest
    ]
    return {"chunks": [pub(chunk) for chunk in S["chunks"]], "mastery": mastery,
            "history": S["history"][-30:], "plan": plan, "llm": bool(client)}


@app.post("/api/upload")
async def upload(files: list[UploadFile] = File(...)):
    for uploaded_file in files:
        data = await uploaded_file.read()
        extension = Path(uploaded_file.filename or "").suffix.lower()
        if extension == ".pdf":
            items, item_type = parse_pdf(data)
        elif extension == ".pptx":
            items, item_type = parse_pptx(data)
        elif extension in (".srt", ".vtt", ".txt", ".md"):
            items, item_type = parse_lecture(data.decode("utf-8", "ignore"))
        elif extension in MEDIA:
            items, item_type = parse_media(data, uploaded_file.filename or "upload")
        else:
            raise HTTPException(400, f"Unsupported file type: {extension or uploaded_file.filename}")
        if not items:
            raise HTTPException(400, f"No readable text found in {uploaded_file.filename}.")
        add(items, item_type, uploaded_file.filename or "upload")
    reindex()
    save()
    return state()


SAMPLE = [
    ("Lecture", "04:12", "A neural network makes a prediction, then a loss function measures how wrong it was. Training means changing the weights to make that loss smaller."),
    ("Lecture", "07:40", "Gradient descent computes the gradient of the loss with respect to each weight, then steps each weight a small amount in the opposite direction. The step size is the learning rate."),
    ("Lecture", "15:05", "If the learning rate is too large the loss overshoots and may diverge. If it is too small, training crawls or stalls in a flat region."),
    ("Textbook", "p.141", "Backpropagation applies the chain rule layer by layer from the output back to the input, reusing intermediate results so all gradients are computed in one backward pass."),
    ("Textbook", "p.148", "Stochastic gradient descent estimates the gradient from a small random mini-batch rather than the full dataset, trading exact gradients for much cheaper steps."),
    ("Textbook", "p.153", "Overfitting occurs when a model memorizes training data and fails on new data. Regularization, dropout and early stopping on a validation set reduce it."),
    ("Slide", "slide 9", "Activation functions such as ReLU and sigmoid add non-linearity. Without them, stacked layers collapse into a single linear map."),
    ("Slide", "slide 12", "An epoch is one pass over the full training set. A batch is the set of examples used for one weight update."),
    ("Slide", "slide 17", "Validation loss rising while training loss keeps falling is the signal that the model is overfitting."),
]


@app.post("/api/sample")
def sample():
    S.update(chunks=[], mastery={}, history=[])
    for item_type, reference, text in SAMPLE:
        add([(reference, text)], item_type, "sample-course")
    reindex()
    save()
    return state()


@app.post("/api/reset")
def reset():
    S.update(chunks=[], mastery={}, history=[])
    reindex()
    save()
    return state()


class Ask(BaseModel):
    q: str


@app.post("/api/ask")
def ask(request: Ask):
    hits = retrieve(request.q)
    if not hits:
        return {"answer": "Nothing in your materials matches that question.", "sources": [], "mode": "none"}
    if client:
        context = "\n".join(f"[{i + 1}] ({hit['type']} {hit['ref']}) {hit['text']}" for i, hit in enumerate(hits))
        answer = llm(
            f"Sources:\n{context}\n\nQuestion: {request.q}",
            system="You are a study tutor. Answer using ONLY the numbered sources. Cite every claim as [n]. If the sources do not answer the question, say so. Max 120 words.",
        )
        mode = "llm"
    else:
        answer = " ".join(f"{hit['text']} [{i + 1}]" for i, hit in enumerate(hits[:2]))
        mode = "extractive"
    return {"answer": answer, "sources": [pub(hit) for hit in hits], "mode": mode}


def cloze(source):
    if vec is None or mat is None:
        raise HTTPException(400, "Upload materials or load the sample course first.")
    terms = vec.get_feature_names_out()
    row_index = S["chunks"].index(source)
    words = set(re.findall(r"[a-z]+", source["text"].lower()))
    row = mat[row_index].toarray().ravel()
    candidates = [terms[j] for j in row.argsort()[::-1] if " " not in terms[j] and terms[j] in words]
    if not candidates:
        raise HTTPException(400, "Could not generate a quiz question from this passage.")
    answer = candidates[0]
    sentence = next((s for s in re.split(r"(?<=[.!?])\s+", source["text"]) if re.search(rf"\b{re.escape(answer)}\b", s, re.I)), source["text"])
    other = [terms[j] for j in mat.sum(axis=0).A1.argsort()[::-1] if " " not in terms[j] and terms[j] not in words][:30]
    options = [answer] + random.sample(other, min(3, len(other)))
    random.shuffle(options)
    question = re.sub(rf"\b{re.escape(answer)}\b", "_____", sentence, flags=re.I)
    return {"q": "Fill in the blank: " + question, "options": options, "answer": options.index(answer)}


def gen_question(source, hard):
    prompt = (
        f"Passage ({source['type']} {source['ref']}):\n{source['text']}\n\n"
        f"Write ONE {'harder, application-style' if hard else 'foundational'} multiple-choice question answerable only from this passage. "
        'Reply with JSON only: {"q": str, "options": [3 or 4 strings], "answer": index of correct option}'
    )
    try:
        raw = llm(prompt, max_tokens=400)
        match = re.search(r"\{.*\}", raw, re.S)
        if not match:
            return None
        result = json.loads(match.group(0))
        if 2 <= len(result["options"]) <= 4 and 0 <= int(result["answer"]) < len(result["options"]):
            return {"q": result["q"], "options": result["options"], "answer": int(result["answer"])}
    except Exception:
        pass
    return None


class QReq(BaseModel):
    concept: str | None = None


@app.post("/api/quiz/next")
def quiz_next(request: QReq):
    if not S["mastery"]:
        raise HTTPException(400, "Upload materials or load the sample course first.")
    concept = request.concept if request.concept in S["mastery"] else min(S["mastery"], key=S["mastery"].get)
    hard = S["mastery"][concept] >= 0.5
    candidates = [chunk for chunk in S["chunks"] if chunk.get("concept") == concept]
    if not candidates:
        raise HTTPException(400, "No source material is available for this concept.")
    source = random.choice(candidates)
    question = (gen_question(source, hard) if client else None) or cloze(source)
    return {**question, "concept": concept, "difficulty": 2 if hard else 1,
            "source": pub(source), "mode": "llm" if client and question else "cloze"}


class Ans(BaseModel):
    concept: str
    correct: bool
    difficulty: int


@app.post("/api/quiz/answer")
def quiz_answer(request: Ans):
    mastery = S["mastery"].get(request.concept, 0.3)
    S["mastery"][request.concept] = (
        mastery + (1 - mastery) * (0.45 if request.difficulty == 2 else 0.35)
        if request.correct else mastery * 0.7
    )
    average = sum(S["mastery"].values()) / len(S["mastery"]) if S["mastery"] else 0
    S["history"].append({"t": int(time.time()), "avg": average})
    save()
    return state()


reindex()

# Fix for Vercel: only mount static files when the folder is present in the deployment.
STATIC_DIR = ROOT / "static"
if STATIC_DIR.is_dir():
    app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")
